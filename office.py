"""LibreOffice discovery and isolated PDF export helpers."""

from __future__ import annotations

import os
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
import threading
import zipfile

from openpyxl import load_workbook


class OfficeError(RuntimeError):
    """Raised when LibreOffice cannot be found or cannot export a workbook."""


APP_DIR = Path(__file__).resolve().parent
PORTABLE_SOFFICE = APP_DIR / "runtime" / "libreoffice" / "libreoffice" / "LibreOfficeDev.app" / "Contents" / "MacOS" / "soffice"
PORTABLE_FONTCONFIG_DIR = (PORTABLE_SOFFICE.parent / "../Resources/fontconfig").resolve()
PORTABLE_FONTCONFIG_FILE = PORTABLE_FONTCONFIG_DIR / "fonts.conf"
BUNDLED_SOFFICE = Path(
    "/Users/koko/.cache/codex-runtimes/codex-primary-runtime/"
    "dependencies/bin/override/soffice"
)
BUNDLED_FONTCONFIG_DIR = (
    BUNDLED_SOFFICE.parent
    / "../../native/libreoffice-headless/libreoffice/LibreOfficeDev.app/Contents/Resources/fontconfig"
).resolve()
BUNDLED_FONTCONFIG_FILE = BUNDLED_FONTCONFIG_DIR / "fonts.conf"
KNOWN_SOFFICE_PATHS = (
    Path("/Applications/LibreOffice.app/Contents/MacOS/soffice"),
    Path("/Applications/LibreOffice.app/Contents/MacOS/soffice.bin"),
    Path("/usr/local/bin/soffice"),
    Path("/opt/homebrew/bin/soffice"),
    Path("/usr/bin/soffice"),
)
EXPORT_TIMEOUT_SECONDS = 120
PDF_DESTINATION_LOCK = threading.Lock()
UPPER_DIGITS = "零壹贰叁肆伍陆柒捌玖"
SMALL_UNITS = ("", "拾", "佰", "仟")
BIG_UNITS = ("", "万", "亿", "兆")


def _is_executable(path: Path) -> bool:
    return path.is_file() and os.access(path, os.X_OK)


def find_soffice(configured_path: str) -> Path:
    """Resolve a usable ``soffice`` executable from configuration or the host."""
    configured = str(configured_path or "").strip()
    if configured:
        path = Path(configured).expanduser()
        if not _is_executable(path):
            reason = "does not exist" if not path.exists() else "not executable"
            raise OfficeError(f"configured soffice {path} {reason}")
        return path

    candidates: list[Path] = [PORTABLE_SOFFICE, BUNDLED_SOFFICE]
    path_candidate = shutil.which("soffice")
    if path_candidate:
        candidates.append(Path(path_candidate))
    candidates.extend(KNOWN_SOFFICE_PATHS)
    seen: set[Path] = set()
    for path in candidates:
        path = Path(path).expanduser()
        if path in seen:
            continue
        seen.add(path)
        if _is_executable(path):
            return path
    raise OfficeError("soffice executable not found")


def _remove_directory(path: Path | None) -> None:
    if path is None:
        return
    try:
        shutil.rmtree(path, ignore_errors=True)
    except Exception:
        # Cleanup must never replace the conversion error.
        pass


def _conversion_environment(soffice_path: Path, profile_dir: Path) -> dict[str, str]:
    """Configure bundled LibreOffice to discover macOS fonts during headless export."""
    environment = os.environ.copy()
    try:
        resolved = Path(soffice_path).resolve()
        is_portable = resolved == PORTABLE_SOFFICE.resolve()
        is_bundled = resolved == BUNDLED_SOFFICE.resolve()
    except OSError:
        is_portable = is_bundled = False
    fontconfig_file = PORTABLE_FONTCONFIG_FILE if is_portable else BUNDLED_FONTCONFIG_FILE
    fontconfig_dir = PORTABLE_FONTCONFIG_DIR if is_portable else BUNDLED_FONTCONFIG_DIR
    if (is_portable or is_bundled) and fontconfig_file.is_file():
        font_cache_dir = profile_dir / "fontconfig-cache"
        font_cache_dir.mkdir(parents=True, exist_ok=True)
        environment.update({
            "FONTCONFIG_FILE": str(fontconfig_file),
            "FONTCONFIG_PATH": str(fontconfig_dir),
            "XDG_CACHE_HOME": str(font_cache_dir),
        })
    return environment


def _as_decimal(value) -> Decimal:
    """Convert a numeric workbook value to Decimal, treating blanks as zero."""
    if value in (None, ""):
        return Decimal("0")
    try:
        return Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError):
        return Decimal("0")


def _upper_integer(value: int) -> str:
    if value == 0:
        return UPPER_DIGITS[0]
    if value < 0:
        return "负" + _upper_integer(-value)

    def group_text(group: int) -> str:
        result: list[str] = []
        zero_pending = False
        for position in range(3, -1, -1):
            digit = (group // (10 ** position)) % 10
            if digit:
                if zero_pending and result:
                    result.append(UPPER_DIGITS[0])
                result.append(UPPER_DIGITS[digit] + SMALL_UNITS[position])
                zero_pending = False
            elif result:
                zero_pending = True
        return "".join(result)

    groups: list[int] = []
    remaining = value
    while remaining:
        groups.append(remaining % 10000)
        remaining //= 10000

    result: list[str] = []
    zero_pending = False
    for index in range(len(groups) - 1, -1, -1):
        group = groups[index]
        if not group:
            if result:
                zero_pending = True
            continue
        if result and (zero_pending or group < 1000):
            result.append(UPPER_DIGITS[0])
        result.append(group_text(group))
        if index:
            result.append(BIG_UNITS[index] if index < len(BIG_UNITS) else "")
        zero_pending = False
    return "".join(result)


def _amount_to_upper(value) -> str:
    """Format a non-negative workbook total like the template's DBNum2 formula."""
    amount = _as_decimal(value).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    negative = amount < 0
    amount = abs(amount)
    integer = int(amount)
    cents = int((amount - integer) * 100)
    jiao, fen = divmod(cents, 10)

    result = _upper_integer(integer) + "元"
    if not jiao and not fen:
        result += "整"
    elif not jiao:
        if integer:
            result += "零"
        result += UPPER_DIGITS[fen] + "分"
    elif not fen:
        result += UPPER_DIGITS[jiao] + "角"
    else:
        result += UPPER_DIGITS[jiao] + "角" + UPPER_DIGITS[fen] + "分"
    return ("负" if negative else "") + result


def _calculate_template_total(ws) -> Decimal:
    """Evaluate the compact template's J20 total from its input cells."""
    allowance_formula = ws["K7"].value
    allowance = _as_decimal(allowance_formula)
    if isinstance(allowance_formula, str):
        match = re.search(r"J7\s*\*\s*([+-]?(?:\d+(?:\.\d*)?|\.\d+))", allowance_formula, re.IGNORECASE)
        if match:
            allowance = _as_decimal(ws["J7"].value) * _as_decimal(match.group(1))

    total = allowance
    for row_number in range(9, 20):
        total += sum((_as_decimal(ws[f"{column}{row_number}"].value) for column in ("E", "F", "H", "I")), Decimal("0"))
    return total


def _prepare_pdf_workbook(workbook_path: Path, conversion_dir: Path) -> Path:
    """Create a PDF-only copy with a LibreOffice-compatible uppercase total."""
    try:
        workbook = load_workbook(workbook_path)
    except (OSError, ValueError, zipfile.BadZipFile):
        # Unit-test fakes and non-Excel inputs should still reach the converter.
        return workbook_path

    try:
        worksheet = next(
            (sheet for sheet in workbook.worksheets if sheet.title in {"差旅报销单", "1月"}),
            None,
        )
        formula = worksheet["C20"].value if worksheet is not None else None
        if worksheet is None or not (isinstance(formula, str) and "[dbnum2]" in formula.lower()):
            return workbook_path

        worksheet["C20"] = _amount_to_upper(_calculate_template_total(worksheet))
        prepared_path = conversion_dir / workbook_path.name
        workbook.save(prepared_path)
        return prepared_path
    finally:
        workbook.close()


def _reserve_pdf_destination(output_dir: Path, stem: str) -> Path:
    candidate = output_dir / f"{stem}.pdf"
    suffix = 2
    while True:
        try:
            fd = os.open(candidate, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
        except FileExistsError:
            candidate = output_dir / f"{stem}-{suffix}.pdf"
            suffix += 1
            continue
        except OSError as exc:
            raise OfficeError(f"could not reserve PDF destination: {exc}") from exc
        try:
            os.close(fd)
        except OSError as exc:
            try:
                candidate.unlink(missing_ok=True)
            except OSError:
                pass
            raise OfficeError(f"could not reserve PDF destination: {exc}") from exc
        return candidate


def export_pdf(workbook_path: Path, output_dir: Path, soffice_path: Path) -> Path:
    """Recalculate and export one workbook through an isolated LibreOffice profile."""
    workbook_path = Path(workbook_path)
    output_dir = Path(output_dir)
    soffice_path = Path(soffice_path)
    if not workbook_path.is_file():
        raise OfficeError(f"workbook does not exist: {workbook_path}")
    if not _is_executable(soffice_path):
        raise OfficeError(f"soffice is not executable: {soffice_path}")

    profile_dir: Path | None = None
    conversion_dir: Path | None = None
    reserved_destination: Path | None = None
    try:
        try:
            output_dir.mkdir(parents=True, exist_ok=True)
            profile_dir = Path(tempfile.mkdtemp(prefix=".soffice-profile-", dir=output_dir))
            conversion_dir = Path(tempfile.mkdtemp(prefix=".pdf-convert-", dir=output_dir))
        except OSError as exc:
            raise OfficeError(f"could not prepare PDF conversion: {exc}") from exc
        pdf_workbook_path = _prepare_pdf_workbook(workbook_path, conversion_dir)
        command = [
            str(soffice_path),
            "--headless",
            f"-env:UserInstallation={profile_dir.resolve().as_uri()}",
            "--convert-to",
            "pdf",
            "--outdir",
            str(conversion_dir),
            str(pdf_workbook_path),
        ]
        try:
            environment = _conversion_environment(soffice_path, profile_dir)
        except OSError as exc:
            raise OfficeError(f"could not prepare PDF conversion: {exc}") from exc
        try:
            completed = subprocess.run(
                command,
                capture_output=True,
                text=True,
                timeout=EXPORT_TIMEOUT_SECONDS,
                check=False,
                env=environment,
            )
        except subprocess.TimeoutExpired as exc:
            stderr = exc.stderr.decode(errors="replace") if isinstance(exc.stderr, bytes) else (exc.stderr or "")
            detail = f": {stderr.strip()}" if stderr.strip() else ""
            raise OfficeError(f"soffice conversion timed out{detail}") from exc
        except OSError as exc:
            raise OfficeError(f"could not run soffice: {exc}") from exc

        stderr = (completed.stderr or "").strip()
        if completed.returncode != 0:
            detail = f": {stderr}" if stderr else ""
            raise OfficeError(f"soffice conversion failed (exit {completed.returncode}){detail}")

        converted_pdf = conversion_dir / f"{pdf_workbook_path.stem}.pdf"
        if not converted_pdf.is_file():
            detail = f": {stderr}" if stderr else ""
            raise OfficeError(f"soffice did not produce a PDF{detail}")
        with PDF_DESTINATION_LOCK:
            reserved_destination = _reserve_pdf_destination(output_dir, workbook_path.stem)
            try:
                os.replace(converted_pdf, reserved_destination)
            except OSError as exc:
                try:
                    reserved_destination.unlink(missing_ok=True)
                except OSError:
                    pass
                raise OfficeError(f"could not move generated PDF: {exc}") from exc
            destination = reserved_destination
            reserved_destination = None
        return destination
    finally:
        if reserved_destination is not None:
            try:
                reserved_destination.unlink(missing_ok=True)
            except OSError:
                pass
        _remove_directory(profile_dir)
        _remove_directory(conversion_dir)
