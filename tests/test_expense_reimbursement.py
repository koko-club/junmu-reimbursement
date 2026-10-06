from pathlib import Path
import io
import json
import shutil
import tempfile
import unittest

from openpyxl import load_workbook
from pypdf import PdfReader

from generator import generate_workbook
from validation import ValidationError, validate_payload
from office import OfficeError, _prepare_pdf_workbook, export_pdf, find_soffice
from tests.http_helpers import RunningApp
from tests.test_web_reimbursements import multipart


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "resources" / "费用报销单模板.xlsx"


def expense_payload(**changes):
    value = {
        "form_type": "expense",
        "date": "2026-09-30",
        "department": "技术部",
        "traveler": "张三",
        "rows": [
            {"project": "办公用品", "summary": "购买纸张", "amount": "123.45"},
            {},
            {"project": "交通费", "summary": "市内打车", "amount": "6.05"},
        ],
    }
    value.update(changes)
    return value


class ExpenseValidationTest(unittest.TestCase):
    def test_normalizes_eleven_rows_and_ignores_untrusted_profile(self):
        result = validate_payload(expense_payload(), traveler="李四", department="财务部")
        self.assertEqual(result["form_type"], "expense")
        self.assertEqual((result["traveler"], result["department"]), ("李四", "财务部"))
        self.assertEqual(len(result["rows"]), 11)
        self.assertEqual(result["rows"][0]["amount"], 123.45)
        self.assertEqual(result["rows"][1], {})

    def test_requires_date_and_complete_line(self):
        for changes in (
            {"date": ""},
            {"rows": [{"project": "办公用品", "amount": 1}]},
            {"rows": [{"summary": "纸张", "amount": 1}]},
            {"rows": [{"project": "办公用品", "summary": "纸张"}]},
            {"rows": [{"project": "办公用品", "summary": "纸张", "amount": -1}]},
            {"rows": [{}] * 12},
            {"rows": [{}] * 11},
        ):
            with self.subTest(changes=changes), self.assertRaises(ValidationError):
                validate_payload(expense_payload(**changes))


class ExpenseWorkbookTest(unittest.TestCase):
    def test_maps_rows_and_formulas_to_supplied_template(self):
        self.assertTrue(TEMPLATE.is_file())
        with tempfile.TemporaryDirectory() as directory:
            result = generate_workbook(TEMPLATE, Path(directory), validate_payload(expense_payload()), [])
            workbook = load_workbook(result.path)
            try:
                sheet = workbook["费用报销单"]
                self.assertEqual(sheet["A3"].value, "报销日期：2026/09/30")
                self.assertEqual(sheet["B4"].value, "技术部")
                self.assertEqual(sheet["A7"].value, "办公用品")
                self.assertEqual(sheet["D7"].value, "购买纸张")
                self.assertEqual(sheet["H7"].value, 123.45)
                self.assertIsNone(sheet["H8"].value)
                self.assertEqual(sheet["H9"].value, 6.05)
                self.assertIsNone(sheet["H17"].value)
                self.assertEqual(sheet["H18"].value, "=SUM(H7:H17)")
                self.assertIn("H18", sheet["C20"].value)
                self.assertEqual(sheet.row_dimensions[7].height, 15)
                self.assertEqual(sheet.print_area, "'费用报销单'!$A$1:$M$23")
            finally:
                workbook.close()

    def test_pdf_copy_contains_calculated_chinese_uppercase(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            result = generate_workbook(TEMPLATE, root, validate_payload(expense_payload()), [])
            conversion = root / "conversion"
            conversion.mkdir()
            prepared_path = _prepare_pdf_workbook(result.path, conversion)
            prepared = load_workbook(prepared_path)
            try:
                self.assertEqual(prepared["费用报销单"]["C20"].value, "壹佰贰拾玖元伍角")
            finally:
                prepared.close()

    def test_rendered_pdf_detail_spacing_matches_travel(self):
        try:
            soffice = find_soffice("")
        except OfficeError:
            self.skipTest("LibreOffice is unavailable")

        markers = [f"ROW{number:02d}" for number in range(1, 12)]
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            expense_rows = [
                {"project": marker, "summary": "内容", "amount": number}
                for number, marker in enumerate(markers, 1)
            ]
            travel_rows = [
                {"date": "2026-09-30", "origin": marker, "destination": "杭州", "public_amount": number}
                for number, marker in enumerate(markers, 1)
            ]
            forms = (
                (TEMPLATE, validate_payload(expense_payload(rows=expense_rows))),
                (ROOT / "resources" / "差旅报销单模板.xlsx", {
                    "date": "2026-09-30", "traveler": "李四", "rows": travel_rows,
                }),
            )
            positions = []
            for template, payload in forms:
                workbook = generate_workbook(template, root, payload, []).path
                pdf = export_pdf(workbook, root, soffice)
                row_positions = {}

                def collect(text, cm, tm, _font, _size):
                    marker = text.strip()
                    if marker in markers:
                        row_positions[marker] = cm[5] + tm[5]

                PdfReader(pdf).pages[0].extract_text(visitor_text=collect)
                self.assertEqual(set(row_positions), set(markers))
                positions.append([row_positions[marker] for marker in markers])

            expense_y, travel_y = positions
            for number in range(10):
                with self.subTest(row=number + 1):
                    self.assertAlmostEqual(
                        expense_y[number] - expense_y[number + 1],
                        travel_y[number] - travel_y[number + 1],
                        delta=0.5,
                    )


class ExpenseFrontendTest(unittest.TestCase):
    def test_switcher_exposes_eleven_expense_rows(self):
        markup = (ROOT / "templates" / "index.html").read_text(encoding="utf-8")
        script = (ROOT / "static" / "app.js").read_text(encoding="utf-8")
        self.assertIn('id="expense-rows"', markup)
        self.assertEqual(markup.count('class="expense-row"'), 11)
        self.assertIn('id="expense-total"', markup)
        self.assertIn('id="expense-total-upper"', markup)
        self.assertNotIn('expense-development-dialog', markup)
        self.assertIn("selectType('expense')", script)


class ExpenseHttpFlowTest(unittest.TestCase):
    def test_generate_pdf_and_edit_same_expense_record(self):
        try:
            find_soffice("")
        except OfficeError:
            self.skipTest("LibreOffice is unavailable")
        with RunningApp() as running:
            shutil.copyfile(TEMPLATE, running.root / TEMPLATE.name)
            self.assertEqual(running.client.post_json("/api/setup", {
                "username": "admin", "password": "administrator1",
                "real_name": "管理员", "department": "管理部",
            }).status, 201)
            admin = running.new_client()
            self.assertEqual(admin.post_json("/api/login", {
                "username": "admin", "password": "administrator1",
            }).status, 200)
            self.assertEqual(running.client.post_json("/api/register", {
                "username": "alice", "password": "correct horse battery staple",
                "real_name": "张三", "department": "技术部",
            }).status, 201)
            with running.server.database.connect() as connection:
                ids = {row["username"]: row["id"] for row in connection.execute("SELECT id, username FROM users")}
            running.users.approve(ids["admin"], ids["alice"])
            client = running.new_client()
            self.assertEqual(client.post_json("/api/login", {
                "username": "alice", "password": "correct horse battery staple",
            }).status, 200)

            def submit(path, payload):
                body, content_type = multipart([("payload", json.dumps(payload), None, "application/json")])
                return client.request("POST", path, body=body, headers={
                    "Content-Type": content_type,
                    "Content-Length": str(len(body)),
                    "X-CSRF-Token": client.csrf_for(path),
                })

            created = submit("/api/reimbursements/generate", expense_payload())
            self.assertEqual(created.status, 200, created.text)
            record = created.json()["record"]
            self.assertEqual(record["reimbursement_amount"], "129.50")
            pdf = client.get(created.json()["pdf_url"])
            self.assertEqual(pdf.status, 200)
            self.assertIn("壹佰贰拾玖元伍角", PdfReader(io.BytesIO(pdf.body)).pages[0].extract_text())
            edit_path = f'/api/reimbursements/{record["id"]}/edit'
            self.assertEqual(client.get(edit_path).json()["payload"]["form_type"], "expense")
            updated_payload = expense_payload()
            updated_payload["rows"][0]["amount"] = "124.45"
            updated_payload["keep_invoices"] = []
            updated = submit(f'/api/reimbursements/{record["id"]}/regenerate', updated_payload)
            self.assertEqual(updated.status, 200, updated.text)
            self.assertEqual(updated.json()["record"]["id"], record["id"])
            self.assertEqual(updated.json()["record"]["reimbursement_amount"], "130.50")
            self.assertEqual(len(client.get("/api/reimbursements?scope=active").json()["reimbursements"]), 1)
