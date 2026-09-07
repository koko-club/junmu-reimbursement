"""Real generator -> workbook -> PDF verification for the reimbursement flow."""

from __future__ import annotations

import hashlib
from pathlib import Path
import sys
import tempfile
import unittest

from PIL import Image, ImageDraw
from openpyxl import load_workbook

APP_DIR = Path(__file__).resolve().parents[1]
if str(APP_DIR) not in sys.path:
    sys.path.insert(0, str(APP_DIR))

from generator import generate_workbook
from office import OfficeError, export_pdf, find_soffice


TEMPLATE = Path("/Users/koko/Work/3.报销/差旅报销单模版（拷贝备用）_副本.xlsx")
BUNDLED_SOFFICE = Path(
    "/Users/koko/.cache/codex-runtimes/codex-primary-runtime/"
    "dependencies/bin/override/soffice"
)


def complete_payload() -> dict:
    """Return a payload that exercises public transit, self-driving, and blanks."""
    return {
        "date": "2026-09-04",
        "department": "技术部",
        "traveler": "张三",
        "reason": "客户拜访",
        "days": 2,
        "allowance": 77.5,
        "rows": [
            {
                "date": "2026-09-01",
                "origin": "上海",
                "destination": "杭州",
                "transport": "高铁",
                "public_amount": 12.5,
                "receipts": 2,
            },
            {
                "date": "2026-09-02",
                "origin": "杭州",
                "destination": "宁波",
                "transport": "自驾",
                "mileage": 180,
                "toll": 20,
                "lodging": 200,
                "receipts": 1,
            },
            {},
        ],
    }


def make_png(path: Path, label: str, color: tuple[int, int, int]) -> None:
    image = Image.new("RGB", (320, 180), color)
    draw = ImageDraw.Draw(image)
    draw.text((20, 70), label, fill="white")
    image.save(path, format="PNG")


class EndToEndTest(unittest.TestCase):
    def test_complete_generation_and_optional_pdf_conversion(self):
        template_hash_before = hashlib.sha256(TEMPLATE.read_bytes()).hexdigest()
        template_mtime_before = TEMPLATE.stat().st_mtime_ns

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            first_image = root / "route-01.png"
            second_image = root / "route-02.png"
            make_png(first_image, "route 1", (34, 105, 170))
            make_png(second_image, "route 2", (42, 135, 80))
            payload = complete_payload()

            first = generate_workbook(
                TEMPLATE,
                root,
                payload,
                image_paths=[first_image, second_image],
            )
            self.assertTrue(first.path.is_file())
            self.assertEqual(first.screenshot_count, 2)
            self.assertEqual(first.receipt_count, 3)

            source_ws = load_workbook(TEMPLATE, data_only=False, read_only=False)["1月"]
            source_signatures = [source_ws.cell(22, column).value for column in range(1, 13)]
            source_ws.parent.close()

            workbook = load_workbook(first.path, data_only=False, read_only=False)
            try:
                self.assertEqual(workbook.sheetnames, ["差旅报销单", "里程截图01", "里程截图02"])
                ws = workbook["差旅报销单"]
                self.assertEqual(ws.max_row, 62)
                self.assertEqual(ws["A3"].value, "报销日期：2026/09/04")
                self.assertEqual(ws["B4"].value, "技术部")
                self.assertEqual(ws["B5"].value, "张三")
                self.assertEqual(ws["E5"].value, "客户拜访")
                self.assertEqual(ws["J7"].value, 2)
                self.assertEqual(ws["K7"].value, "=J7*77.5")
                self.assertEqual(ws["L5"].value, "附\n单\n据\n\n张")
                self.assertEqual(ws["A25"].value, "详见后附里程截图，共 2 张")

                self.assertEqual(ws["G9"].value, '=IF(F9*1=0,"",F9*1)')
                self.assertEqual(ws["G10"].value, '=IF(F10*1=0,"",F10*1)')
                self.assertEqual(
                    ws["J9"].value,
                    '=IF(F9*1 + E9 + H9 + I9 = 0, "", F9*1 + E9 + H9 + I9)',
                )
                self.assertEqual(
                    ws["J10"].value,
                    '=IF(F10*1 + E10 + H10 + I10 = 0, "", F10*1 + E10 + H10 + I10)',
                )
                # Receipt count is driven by detail receipts, not screenshot count.
                self.assertEqual(ws["K9"].value, 2)
                self.assertEqual(ws["K10"].value, 1)

                # The third detail row remains empty instead of receiving numeric zeroes.
                for column in ("A", "B", "C", "D", "E", "F", "H", "I", "K"):
                    self.assertIsNone(ws[f"{column}11"].value, f"{column}11 should stay blank")
                self.assertEqual(ws["A22"].value, source_signatures[0])
                self.assertEqual(ws["J22"].value, source_signatures[9])
                self.assertEqual(
                    [ws.cell(22, column).value for column in range(1, 13)],
                    source_signatures,
                )
                self.assertEqual(len(workbook["里程截图01"]._images), 1)
                self.assertEqual(len(workbook["里程截图02"]._images), 1)
            finally:
                workbook.close()

            soffice = None
            try:
                soffice = find_soffice(str(BUNDLED_SOFFICE))
            except OfficeError:
                # Generator assertions remain mandatory when PDF conversion is unavailable.
                pass

            first_pdf = None
            if soffice is not None:
                first_pdf = export_pdf(first.path, root, soffice)
                self.assertTrue(first_pdf.is_file())
                self.assertGreater(first_pdf.stat().st_size, 0)
                try:
                    from pypdf import PdfReader
                except ImportError:  # pragma: no cover - bundled runtime currently includes pypdf
                    PdfReader = None
                if PdfReader is not None:
                    reader = PdfReader(str(first_pdf))
                    self.assertGreaterEqual(len(reader.pages), 3)
                    first_text = reader.pages[0].extract_text() or ""
                    self.assertIn("差旅费用报销单", first_text)
                    self.assertIn("2026/09/04", first_text)
                    self.assertIn("张三", first_text)
                    self.assertNotIn("Err:502", first_text)
                    self.assertIn("伍佰陆拾柒元伍角", first_text)
                    appendix_one_text = reader.pages[1].extract_text() or ""
                    appendix_two_text = reader.pages[2].extract_text() or ""
                    self.assertIn("1/2", appendix_one_text)
                    self.assertIn("2/2", appendix_two_text)
                    self.assertTrue(any(image.data for image in reader.pages[1].images))

            # Keep a byte-for-byte snapshot so the collision generation below
            # proves that the first workbook was not overwritten.
            first_bytes = first.path.read_bytes()
            second = generate_workbook(
                TEMPLATE,
                root,
                payload,
                image_paths=[first_image, second_image],
            )
            self.assertTrue(second.path.is_file())
            self.assertEqual(second.path.name, f"{first.path.stem}-2.xlsx")
            self.assertEqual(first.path.read_bytes(), first_bytes)
            if soffice is not None:
                second_pdf = export_pdf(second.path, root, soffice)
                self.assertEqual(second_pdf.name, f"{first_pdf.stem}-2.pdf")
                self.assertTrue(second_pdf.is_file())
                self.assertGreater(second_pdf.stat().st_size, 0)

        self.assertEqual(hashlib.sha256(TEMPLATE.read_bytes()).hexdigest(), template_hash_before)
        self.assertEqual(TEMPLATE.stat().st_mtime_ns, template_mtime_before)


if __name__ == "__main__":
    unittest.main()
