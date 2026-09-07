from pathlib import Path
import sys
import tempfile
from concurrent.futures import ThreadPoolExecutor
import unittest
from unittest.mock import MagicMock, patch

from PIL import Image
from openpyxl import load_workbook

APP_DIR = Path(__file__).resolve().parents[1]
if str(APP_DIR) not in sys.path:
    sys.path.insert(0, str(APP_DIR))

from generator import generate_workbook


TEMPLATE = Path("/Users/koko/Work/3.报销/差旅报销单模版（拷贝备用）_副本.xlsx")


def payload(allowance=50, rows=None):
    return {
        "date": "2026-09-04",
        "department": "技术部",
        "traveler": "张三",
        "reason": "客户拜访",
        "days": 2,
        "allowance": allowance,
        "rows": rows if rows is not None else [
            {
                "date": "2026-09-01",
                "origin": "上海",
                "destination": "杭州",
                "transport": "高铁",
                "public_amount": 12.5,
                "mileage": 10,
                "toll": 20.25,
                "lodging": 180,
                "receipts": 2,
            }
        ] + [{}] * 10,
        "receipt_count": 2,
    }


class GeneratorTest(unittest.TestCase):
    def test_mkstemp_failure_removes_reserved_destination(self):
        with tempfile.TemporaryDirectory() as tmp:
            with patch("generator.tempfile.mkstemp", side_effect=OSError("no temp space")):
                with self.assertRaisesRegex(OSError, "no temp space"):
                    generate_workbook(TEMPLATE, Path(tmp), payload(), image_paths=[])
            self.assertEqual(list(Path(tmp).glob("*.xlsx")), [])

    def test_cleanup_failure_does_not_mask_original_error(self):
        with tempfile.TemporaryDirectory() as tmp:
            with patch("generator.os.replace", side_effect=RuntimeError("replace failed")):
                with patch.object(Path, "unlink", side_effect=OSError("unlink failed")):
                    with self.assertRaisesRegex(RuntimeError, "replace failed"):
                        generate_workbook(TEMPLATE, Path(tmp), payload(), image_paths=[])

    def test_workbook_is_closed_after_generation(self):
        from openpyxl import load_workbook as real_load_workbook

        workbook = real_load_workbook(TEMPLATE)
        workbook.close = MagicMock()
        with tempfile.TemporaryDirectory() as tmp:
            with patch("generator.load_workbook", return_value=workbook):
                generate_workbook(TEMPLATE, Path(tmp), payload(), image_paths=[])
        workbook.close.assert_called_once_with()

    def test_compacts_form_and_writes_header_and_exact_formulas(self):
        with tempfile.TemporaryDirectory() as tmp:
            result = generate_workbook(TEMPLATE, Path(tmp), payload(), image_paths=[])
            self.assertTrue(result["path"].exists())
            wb = load_workbook(result["path"], data_only=False)
            self.assertEqual(wb.sheetnames, ["差旅报销单"])
            ws = wb["差旅报销单"]
            self.assertEqual(ws.max_row, 62)
            self.assertEqual(ws.print_area, "'差旅报销单'!$A$1:$L$62")
            self.assertFalse(ws.sheet_view.showGridLines)
            self.assertAlmostEqual(ws.page_margins.top, 0.1)
            self.assertEqual(ws["A3"].value, "报销日期：2026/09/04")
            self.assertEqual(ws["B4"].value, "技术部")
            self.assertEqual(ws["B5"].value, "张三")
            self.assertEqual(ws["E5"].value, "客户拜访")
            self.assertEqual(ws["J7"].value, 2)
            self.assertEqual(ws["K7"].value, "=J7*50")
            self.assertEqual(ws["L5"].value, "附\n单\n据\n\n张")
            self.assertEqual(ws["A25"].value, "详见后附里程截图，共 0 张")
            self.assertEqual(ws["A22"].value, "审核人:")
            self.assertEqual(ws["J22"].value, "报销人:")
            self.assertEqual(ws["A9"].value, "2026-09-01")
            self.assertEqual(ws["B9"].value, "上海")
            self.assertEqual(ws["C9"].value, "杭州")
            self.assertEqual(ws["D9"].value, "高铁")
            self.assertEqual(ws["E9"].value, 12.5)
            self.assertEqual(ws["F9"].value, 10)
            self.assertEqual(ws["G9"].value, '=IF(F9*1=0,"",F9*1)')
            self.assertEqual(ws["J9"].value, '=IF(F9*1 + E9 + H9 + I9 = 0, "", F9*1 + E9 + H9 + I9)')
            self.assertEqual(ws["K9"].value, 2)
            self.assertIsNone(ws["A10"].value)
            self.assertIsNone(ws["F10"].value)
            self.assertEqual(ws["G10"].value, '=IF(F10*1=0,"",F10*1)')
            self.assertEqual(ws["J10"].value, '=IF(F10*1 + E10 + H10 + I10 = 0, "", F10*1 + E10 + H10 + I10)')
            self.assertEqual(ws["J20"].value, "=SUM(J9:J19,K7)")

    def test_custom_allowance_and_receipt_count_excludes_screenshots(self):
        with tempfile.TemporaryDirectory() as tmp:
            image = Path(tmp) / "route.png"
            Image.new("RGB", (120, 80), "red").save(image)
            result = generate_workbook(TEMPLATE, Path(tmp), payload(allowance=77.5), image_paths=[image])
            wb = load_workbook(result["path"], data_only=False)
            ws = wb["差旅报销单"]
            self.assertEqual(ws["K7"].value, "=J7*77.5")
            self.assertEqual(ws["L5"].value, "附\n单\n据\n\n张")
            self.assertEqual(ws["A25"].value, "详见后附里程截图，共 1 张")
            self.assertEqual(result["screenshot_count"], 1)

    def test_blank_dates_remain_blank_in_generated_form(self):
        rows = [{
            "date": "",
            "origin": "上海",
            "destination": "杭州",
            "public_amount": 12,
        }] + [{}] * 10
        with tempfile.TemporaryDirectory() as tmp:
            test_payload = payload(rows=rows)
            test_payload["date"] = ""
            result = generate_workbook(
                TEMPLATE,
                Path(tmp),
                test_payload,
                image_paths=[],
            )
            wb = load_workbook(result.path, data_only=False)
            try:
                ws = wb["差旅报销单"]
                self.assertEqual(ws["A3"].value, "报销日期：")
                self.assertIsNone(ws["A9"].value)
            finally:
                wb.close()

    def test_screenshot_sheet_is_printable_and_preserves_aspect_ratio(self):
        with tempfile.TemporaryDirectory() as tmp:
            image = Path(tmp) / "tall.png"
            Image.new("RGB", (100, 400), "blue").save(image)
            result = generate_workbook(TEMPLATE, Path(tmp), payload(), image_paths=[image])
            wb = load_workbook(result["path"], data_only=False)
            self.assertEqual(wb.sheetnames, ["差旅报销单", "里程截图01"])
            ws = wb["里程截图01"]
            self.assertEqual(ws["A1"].value, "里程截图 1/1")
            self.assertFalse(ws.sheet_view.showGridLines)
            self.assertEqual(str(ws.page_setup.paperSize), str(ws.PAPERSIZE_A4))
            self.assertEqual(ws.page_setup.orientation, ws.ORIENTATION_PORTRAIT)
            self.assertEqual(ws.page_setup.fitToWidth, 1)
            self.assertEqual(ws.page_setup.fitToHeight, 1)
            self.assertEqual(len(ws._images), 1)
            placed = ws._images[0]
            self.assertAlmostEqual(placed.height / placed.width, 4, delta=0.05)

    def test_multiple_screenshots_create_numbered_appendix_sheets(self):
        with tempfile.TemporaryDirectory() as tmp:
            first = Path(tmp) / "first.png"
            second = Path(tmp) / "second.png"
            Image.new("RGB", (80, 80), "red").save(first)
            Image.new("RGB", (160, 80), "green").save(second)
            result = generate_workbook(TEMPLATE, Path(tmp), payload(), image_paths=[first, second])
            wb = load_workbook(result["path"], data_only=False)
            self.assertEqual(wb.sheetnames, ["差旅报销单", "里程截图01", "里程截图02"])
            self.assertEqual(wb["里程截图01"]["A1"].value, "里程截图 1/2")
            self.assertEqual(wb["里程截图02"]["A1"].value, "里程截图 2/2")
            self.assertEqual(result["screenshot_count"], 2)

    def test_output_name_collision_gets_incrementing_suffix(self):
        with tempfile.TemporaryDirectory() as tmp:
            output_dir = Path(tmp)
            first = generate_workbook(TEMPLATE, output_dir, payload(), image_paths=[])
            second = generate_workbook(TEMPLATE, output_dir, payload(), image_paths=[])
            self.assertNotEqual(first["path"], second["path"])
            self.assertTrue(str(second["path"]).endswith("-2.xlsx"))
            self.assertTrue(first["path"].exists())
            self.assertTrue(second["path"].exists())

    def test_concurrent_generations_reserve_distinct_output_names(self):
        with tempfile.TemporaryDirectory() as tmp:
            output_dir = Path(tmp)
            with ThreadPoolExecutor(max_workers=2) as executor:
                futures = [
                    executor.submit(generate_workbook, TEMPLATE, output_dir, payload(), []),
                    executor.submit(generate_workbook, TEMPLATE, output_dir, payload(), []),
                ]
                results = [future.result() for future in futures]
            paths = {result["path"] for result in results}
            self.assertEqual(len(paths), 2)
            self.assertTrue(all(path.exists() for path in paths))
            self.assertEqual(
                {path.name for path in paths},
                {"2026-09-04-张三-差旅报销单.xlsx", "2026-09-04-张三-差旅报销单-2.xlsx"},
            )


if __name__ == "__main__":
    unittest.main()
