"""Generate compact travel reimbursement workbooks from the supplied template."""

from __future__ import annotations

from dataclasses import dataclass
from copy import copy
import os
from pathlib import Path
import tempfile
from typing import Any

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as SpreadsheetImage
from openpyxl.worksheet.page import PageMargins
from PIL import Image as PillowImage

try:
    from .validation import safe_output_stem
except ImportError:  # pragma: no cover - supports direct imports from app directory
    from validation import safe_output_stem


@dataclass(frozen=True)
class GenerationResult:
    path: Path
    screenshot_count: int
    receipt_count: int

    def __getitem__(self, key: str) -> Any:
        if key in {"xlsx_path", "workbook_path", "output_path"}:
            return self.path
        return getattr(self, key)

    def as_dict(self) -> dict[str, Any]:
        return {
            "path": self.path,
            "screenshot_count": self.screenshot_count,
            "receipt_count": self.receipt_count,
        }


def _first_value(row: dict, *names: str) -> Any:
    for name in names:
        if name in row and row[name] not in (None, ""):
            return row[name]
    return None


def _set_or_clear(cell, value: Any) -> None:
    cell.value = value if value not in (None, "") else None


def populate_main_sheet(ws, payload: dict) -> None:
    """Write the header/detail mapping and factor-1 formulas into the compact form."""
    date_text = str(payload.get("date", ""))
    if len(date_text) == 10 and date_text[4] == "-" and date_text[7] == "-":
        date_text = date_text.replace("-", "/")
    ws["A3"] = f"报销日期：{date_text}"
    ws["B4"] = payload.get("department") or None
    ws["B5"] = payload.get("traveler") or None
    ws["E5"] = payload.get("reason") or None
    ws["J7"] = payload.get("days") if payload.get("days") not in (None, "") else None
    allowance = payload.get("allowance", 0)
    ws["K7"] = f"=J7*{allowance}"

    rows = payload.get("rows") or []
    screenshot_count = int(payload.get("screenshot_count", 0) or 0)
    receipt_count = payload.get("receipt_count")
    if receipt_count is None:
        receipt_count = sum((row.get("receipts", 0) or 0) for row in rows if isinstance(row, dict))
    ws["L5"] = "附\n单\n据\n\n张"
    ws["A25"] = f"详见后附里程截图，共 {screenshot_count} 张"

    for row_number in range(9, 20):
        row = rows[row_number - 9] if row_number - 9 < len(rows) and isinstance(rows[row_number - 9], dict) else {}
        populated = any(value not in (None, "") for value in row.values())
        if populated:
            values = {
                "A": row.get("date"),
                "B": row.get("origin"),
                "C": row.get("destination"),
                "D": row.get("transport"),
                "E": _first_value(row, "public_amount", "public_transport_amount", "amount"),
                "F": _first_value(row, "mileage", "driving_mileage"),
                "H": row.get("toll"),
                "I": row.get("lodging"),
                "K": row.get("receipts"),
            }
            for column, value in values.items():
                _set_or_clear(ws[f"{column}{row_number}"], value)
        else:
            for column in ("A", "B", "C", "D", "E", "F", "H", "I", "K"):
                ws[f"{column}{row_number}"] = None
        ws[f"G{row_number}"] = f'=IF(F{row_number}*1=0,"",F{row_number}*1)'
        ws[f"J{row_number}"] = (
            f'=IF(F{row_number}*1 + E{row_number} + H{row_number} + I{row_number} = 0, '
            f'"", F{row_number}*1 + E{row_number} + H{row_number} + I{row_number})'
        )


def _configure_print(ws, print_area: str | None = None) -> None:
    ws.sheet_view.showGridLines = False
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.1, bottom=0.35, header=0.15, footer=0.15)
    if print_area:
        ws.print_area = print_area


def add_screenshot_sheet(wb, image_path: Path, index: int, total: int) -> None:
    """Add one A4-printable worksheet containing a proportionally scaled screenshot."""
    sheet = wb.create_sheet(f"里程截图{index:02d}")
    _configure_print(sheet, "A1:L48")
    sheet["A1"] = f"里程截图 {index}/{total}"
    title_font = copy(sheet["A1"].font)
    title_font.bold = True
    title_font.sz = 14
    sheet["A1"].font = title_font
    with PillowImage.open(image_path) as source:
        width, height = source.size
    max_width, max_height = 700, 900
    scale = min(max_width / width, max_height / height, 1.0)
    image = SpreadsheetImage(str(image_path))
    image.width = max(1, int(round(width * scale)))
    image.height = max(1, int(round(height * scale)))
    sheet.add_image(image, "A3")


def _trim_to_form(ws) -> None:
    original_last_row = ws.max_row
    # Remove merged ranges that belong to the repeated sections before deleting
    # their cells; unmerge_cells cannot process already-deleted placeholders.
    for merged in list(ws.merged_cells.ranges):
        if merged.min_row > 62 or merged.max_row > 62:
            ws.merged_cells.ranges.remove(merged)
    if original_last_row > 62:
        ws.delete_rows(63, original_last_row - 62)


def _set_calculation_mode(wb) -> None:
    calculation = getattr(wb, "calculation", None)
    if calculation is None:
        return
    for name, value in (("fullCalcOnLoad", True), ("forceFullCalc", True), ("calcMode", "auto")):
        try:
            setattr(calculation, name, value)
        except (AttributeError, TypeError):
            pass


def generate_workbook(
    template_path: Path,
    output_dir: Path,
    payload: dict,
    image_paths: list[Path],
) -> GenerationResult:
    """Create a compact workbook and atomically return its final path and metadata."""
    template_path = Path(template_path)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    wb = None
    candidate = None
    candidate_reserved = False
    temporary_path = None
    committed = False
    try:
        wb = load_workbook(template_path)
        if "1月" not in wb.sheetnames:
            raise ValueError("template does not contain the 1月 worksheet")
        main = wb["1月"]
        for sheet in list(wb.worksheets):
            if sheet is not main:
                wb.remove(sheet)
        main.title = "差旅报销单"
        _trim_to_form(main)
        _configure_print(main, "A1:L62")
        normalized = dict(payload)
        normalized["screenshot_count"] = len(image_paths)
        populate_main_sheet(main, normalized)
        for index, image_path in enumerate(image_paths, 1):
            add_screenshot_sheet(wb, Path(image_path), index, len(image_paths))
        _set_calculation_mode(wb)

        receipt_count = normalized.get("receipt_count")
        if receipt_count is None:
            receipt_count = sum((row.get("receipts", 0) or 0) for row in normalized.get("rows", []) if isinstance(row, dict))
        stem = safe_output_stem(normalized.get("date", ""), normalized.get("traveler", ""))
        candidate = output_dir / f"{stem}.xlsx"
        suffix = 2
        # Reserve the destination atomically so concurrent requests cannot select
        # the same collision-free name and overwrite one another.
        while True:
            if candidate.resolve() == template_path.resolve():
                candidate = output_dir / f"{stem}-{suffix}.xlsx"
                suffix += 1
                continue
            try:
                reservation_fd = os.open(candidate, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
            except FileExistsError:
                candidate = output_dir / f"{stem}-{suffix}.xlsx"
                suffix += 1
                continue
            os.close(reservation_fd)
            candidate_reserved = True
            break
        fd, temporary_name = tempfile.mkstemp(prefix=f".{stem}-", suffix=".xlsx", dir=output_dir)
        os.close(fd)
        temporary_path = Path(temporary_name)
        wb.save(temporary_path)
        os.replace(temporary_path, candidate)
        committed = True
    finally:
        if temporary_path is not None:
            try:
                if temporary_path.exists():
                    temporary_path.unlink()
            except OSError:
                pass
        if candidate_reserved and not committed and candidate is not None:
            try:
                candidate.unlink()
            except OSError:
                pass
        if wb is not None:
            try:
                wb.close()
            except OSError:
                pass
    return GenerationResult(candidate, len(image_paths), int(receipt_count or 0))
